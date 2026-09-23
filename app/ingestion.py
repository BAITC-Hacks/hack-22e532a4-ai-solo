from __future__ import annotations

import hashlib
import re
import shutil
import unicodedata
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from pypdf import PdfReader

from .domain import ParsedDocument, ParsedSpan


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SUPPORTED = {".docx", ".pdf", ".xlsx"}
MAX_BYTES = 25 * 1024 * 1024
MAX_UNPACKED = 100 * 1024 * 1024
CLAUSE_RE = re.compile(r"^\s*((?:\d+\.)+\d*|\d+\.)(?:\s+|$)")
MULTI_CLAUSE_RE = re.compile(r"(?<!\S)(?=(?:\d+\.){2,}\s)")


class IngestionError(ValueError):
    pass


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).lower().replace("ё", "е")
    text = re.sub(r"[^\wа-яәіңғүұқөһ]+", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_filename(name: str) -> str:
    clean = Path(name).name
    clean = re.sub(r"[^\w.()\-а-яА-ЯәіңғүұқөһӘІҢҒҮҰҚӨҺ ]+", "_", clean)
    clean = clean.strip(" .")
    if not clean:
        raise IngestionError("Пустое имя файла")
    return clean[:180]


def validate_file(path: Path) -> None:
    if not path.is_file():
        raise IngestionError("Файл не найден")
    if path.stat().st_size == 0:
        raise IngestionError("Файл пуст")
    if path.stat().st_size > MAX_BYTES:
        raise IngestionError("Файл превышает лимит 25 МБ")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise IngestionError(f"Формат {suffix or 'без расширения'} не поддерживается")
    head = path.read_bytes()[:8]
    if suffix in {".docx", ".xlsx"} and not head.startswith(b"PK"):
        raise IngestionError("Расширение не соответствует OOXML-файлу")
    if suffix == ".pdf" and not head.startswith(b"%PDF"):
        raise IngestionError("Расширение не соответствует PDF-файлу")
    if suffix in {'.docx', '.xlsx'}:
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries)>3000 or sum(e.file_size for e in entries)>MAX_UNPACKED:
                    raise IngestionError('OOXML превышает безопасный лимит: 3000 частей / 100 МБ после распаковки')
                if any(e.file_size>2*1024*1024 and e.file_size>max(1,e.compress_size)*200 for e in entries):
                    raise IngestionError('Слишком высокая степень сжатия OOXML')
        except zipfile.BadZipFile as exc:
            raise IngestionError('OOXML-контейнер повреждён') from exc


def store_upload(source: Path, upload_root: Path, comparison_id: str, side: str) -> Path:
    validate_file(source)
    target_dir = (upload_root / comparison_id / side).resolve()
    root = upload_root.resolve()
    if root not in target_dir.parents:
        raise IngestionError("Некорректный путь хранения")
    target_dir.mkdir(parents=True, exist_ok=True)
    digest = sha256_file(source)
    target = target_dir / f"{digest[:12]}_{safe_filename(source.name)}"
    if not target.exists():
        shutil.copy2(source, target)
    return target


def infer_revision(filename: str, spans: list[ParsedSpan]) -> str:
    patterns = [
        re.compile(r"редакци[яи]_?\s*(\d+)", re.IGNORECASE),
        re.compile(r"\bред\.?\s*(\d+)\b", re.IGNORECASE),
    ]
    sample = filename + " " + " ".join(span.text for span in spans[:30])
    for pattern in patterns:
        match = pattern.search(sample)
        if match:
            return f"редакция {match.group(1)}"
    return "версия не указана"


def _text_of(element: ET.Element) -> str:
    parts: list[str] = []
    for node in element.iter():
        if node.tag == W + "t" and node.text:
            parts.append(node.text)
        elif node.tag == W + "tab":
            parts.append("\t")
        elif node.tag in {W + "br", W + "cr"}:
            parts.append(" ")
    return re.sub(r"\s+", " ", "".join(parts)).strip()


def _alpha(number: int) -> str:
    return chr(ord("а") + max(0, min(number - 1, 31)))


def _roman(number: int) -> str:
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for value, glyph in vals:
        while number >= value:
            out += glyph
            number -= value
    return out


def _numbering(zipf: zipfile.ZipFile) -> tuple[dict[str, str], dict[tuple[str, str], dict[str, str]]]:
    try:
        root = ET.fromstring(zipf.read("word/numbering.xml"))
    except KeyError:
        return {}, {}
    abstracts: dict[tuple[str, str], dict[str, str]] = {}
    for abstract in root.findall(W + "abstractNum"):
        aid = abstract.get(W + "abstractNumId", "")
        for level in abstract.findall(W + "lvl"):
            ilvl = level.get(W + "ilvl", "0")
            start = level.find(W + "start")
            fmt = level.find(W + "numFmt")
            text = level.find(W + "lvlText")
            abstracts[(aid, ilvl)] = {
                "start": start.get(W + "val", "1") if start is not None else "1",
                "fmt": fmt.get(W + "val", "decimal") if fmt is not None else "decimal",
                "text": text.get(W + "val", f"%{int(ilvl)+1}.") if text is not None else f"%{int(ilvl)+1}.",
            }
    num_to_abstract: dict[str, str] = {}
    for num in root.findall(W + "num"):
        num_id = num.get(W + "numId", "")
        abstract = num.find(W + "abstractNumId")
        if abstract is not None:
            original = abstract.get(W + "val", "")
            aid = original + ':' + num_id
            for (key, level), spec in list(abstracts.items()):
                if key == original:
                    abstracts[(aid, level)] = dict(spec)
            for override in num.findall(W + 'lvlOverride'):
                level = override.get(W + 'ilvl', '0')
                spec = abstracts.get((aid, level), {})
                custom = override.find(W + 'lvl')
                if custom is not None:
                    for tag, key in [('start','start'),('numFmt','fmt'),('lvlText','text')]:
                        node = custom.find(W + tag)
                        if node is not None:
                            spec[key] = node.get(W + 'val')
                start = override.find(W + 'startOverride')
                if start is not None:
                    spec['start'] = start.get(W + 'val', '1')
                abstracts[(aid, level)] = spec
            num_to_abstract[num_id] = aid
    return num_to_abstract, abstracts


def _paragraph_number(
    paragraph: ET.Element,
    num_to_abstract: dict[str, str],
    levels: dict[tuple[str, str], dict[str, str]],
    counters: dict[str, list[int]],
) -> str | None:
    ppr = paragraph.find(W + "pPr")
    numpr = ppr.find(W + "numPr") if ppr is not None else None
    if numpr is None:
        return None
    num_id_node = numpr.find(W + "numId")
    ilvl_node = numpr.find(W + "ilvl")
    if num_id_node is None:
        return None
    num_id = num_id_node.get(W + "val", "")
    ilvl = int(ilvl_node.get(W + "val", "0")) if ilvl_node is not None else 0
    abstract = num_to_abstract.get(num_id)
    spec = levels.get((abstract or "", str(ilvl)))
    if not spec:
        return None
    if ilvl not in range(9):
        raise IngestionError('Неподдерживаемый уровень нумерации DOCX')
    values = counters.setdefault(num_id, [int(levels.get((abstract or '', str(i)), {}).get('start', '1'))-1 for i in range(9)])
    values[ilvl] += 1
    for idx in range(ilvl + 1, len(values)):
        values[idx] = int(levels.get((abstract or '', str(idx)), {}).get('start', '1'))-1
    template = spec["text"]
    for idx in range(9):
        level_spec = levels.get((abstract or "", str(idx)), {"fmt": "decimal"})
        value = values[idx] if values[idx] >= int(level_spec.get('start','1')) else int(level_spec.get('start','1'))
        fmt = level_spec.get("fmt")
        if fmt not in {'decimal','lowerLetter','lowerRoman','bullet'}:
            return None
        rendered = _alpha(value) if fmt == "lowerLetter" else _roman(value) if fmt == "lowerRoman" else str(value)
        template = template.replace(f"%{idx+1}", rendered)
    return template


def _split_clauses(text: str) -> list[str]:
    pieces = [piece.strip() for piece in MULTI_CLAUSE_RE.split(text) if piece.strip()]
    return pieces or [text]


def parse_docx(path: Path) -> tuple[list[ParsedSpan], list[str]]:
    spans: list[ParsedSpan] = []
    warnings: list[str] = []
    try:
        with zipfile.ZipFile(path) as zipf:
            if "word/document.xml" not in zipf.namelist():
                raise IngestionError("В DOCX отсутствует word/document.xml")
            num_to_abstract, levels = _numbering(zipf)
            styles = {}
            if 'word/styles.xml' in zipf.namelist():
                for style in ET.fromstring(zipf.read('word/styles.xml')).findall(W+'style'):
                    styles[style.get(W+'styleId')] = style
            counters: dict[str, list[int]] = {}
            root = ET.fromstring(zipf.read("word/document.xml"))
            body = root.find(W + "body")
            if body is None:
                raise IngestionError("В DOCX отсутствует тело документа")
            ordinal = 0
            table_no = 0
            for child in body:
                if child.tag == W + "p":
                    text = _text_of(child)
                    if not text:
                        continue
                    ppr = child.find(W+'pPr')
                    style_node = ppr.find(W+'pStyle') if ppr is not None else None
                    if ppr is not None and ppr.find(W+'numPr') is None and style_node is not None:
                        style_id, visited = style_node.get(W+'val'), set()
                        while style_id in styles and style_id not in visited:
                            visited.add(style_id)
                            style = styles[style_id]
                            numpr = style.find(W+'pPr/'+W+'numPr')
                            if numpr is not None:
                                ppr.append(ET.fromstring(ET.tostring(numpr)))
                                break
                            based = style.find(W+'basedOn')
                            style_id = based.get(W+'val') if based is not None else None
                    automatic = _paragraph_number(child, num_to_abstract, levels, counters)
                    if automatic is None and ppr is not None and ppr.find(W+'numPr') is not None:
                        warnings.append('Нумерация одного из абзацев не разрешена; используйте локатор абзаца')
                    if automatic and not CLAUSE_RE.match(text):
                        text = f"{automatic} {text}"
                    for piece in _split_clauses(text):
                        ordinal += 1
                        match = CLAUSE_RE.match(piece)
                        label = match.group(1).rstrip(".") if match else None
                        spans.append(
                            ParsedSpan(ordinal, f"абзац {ordinal}", label, piece, normalize(piece))
                        )
                elif child.tag == W + "tbl":
                    table_no += 1
                    for row_no, row in enumerate(child.findall(W + "tr"), 1):
                        cells = [_text_of(cell) for cell in row.findall(W + "tc")]
                        text = " | ".join(cells)
                        if any(cells):
                            ordinal += 1
                            spans.append(
                                ParsedSpan(
                                    ordinal,
                                    f"таблица {table_no}, строка {row_no}",
                                    None,
                                    text,
                                    normalize(text),
                                    "table_row",
                                )
                            )
    except zipfile.BadZipFile as exc:
        raise IngestionError("DOCX повреждён или не является ZIP-контейнером") from exc
    except (ET.ParseError, ValueError, IndexError) as exc:
        raise IngestionError('DOCX содержит повреждённый XML или некорректную нумерацию') from exc
    if not spans:
        warnings.append("Текст не извлечён; возможно, документ содержит только изображения")
    return spans, warnings


def parse_pdf(path: Path) -> tuple[list[ParsedSpan], list[str]]:
    spans: list[ParsedSpan] = []
    warnings: list[str] = []
    try:
        reader = PdfReader(str(path))
        if len(reader.pages)>500:
            raise IngestionError('PDF превышает лимит 500 страниц')
        for page_no, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if not text:
                warnings.append(f"Страница {page_no}: текст не извлечён, возможно требуется OCR")
                continue
            for block_no, block in enumerate(re.split(r"\n\s*\n|(?<=\.)\s*\n", text), 1):
                block = re.sub(r"\s+", " ", block).strip()
                if not block:
                    continue
                ordinal = len(spans) + 1
                match = CLAUSE_RE.match(block)
                spans.append(
                    ParsedSpan(
                        ordinal,
                        f"страница {page_no}, фрагмент {block_no}",
                        match.group(1).rstrip(".") if match else None,
                        block,
                        normalize(block),
                        "pdf_block",
                    )
                )
    except Exception as exc:
        raise IngestionError(f"Не удалось прочитать PDF: {exc}") from exc
    if not spans:
        warnings.append("PDF не содержит извлекаемого текста; анализ заблокирован до OCR")
    return spans, warnings


def parse_xlsx(path: Path) -> tuple[list[ParsedSpan], list[str]]:
    spans: list[ParsedSpan] = []
    warnings: list[str] = []
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0)>10000 or (sheet.max_column or 0)>100:
                raise IngestionError('XLSX превышает лимит 10000 строк / 100 столбцов на лист')
            for row in sheet.iter_rows():
                if any(cell.data_type=='f' for cell in row):
                    warnings.append(f'Лист {sheet.title}: формулы не вычисляются; экспортируйте значения для полного анализа')
                values = [str(cell.value).strip() if cell.value is not None and cell.data_type!='f' else '' for cell in row]
                if not any(values):
                    continue
                text = " | ".join(values)
                ordinal = len(spans) + 1
                first = next(cell for cell in row if cell.value not in (None, ""))
                last = next(cell for cell in reversed(row) if cell.value not in (None, ""))
                locator = f"лист {sheet.title}, {first.coordinate}:{last.coordinate}"
                match = CLAUSE_RE.match(text)
                spans.append(
                    ParsedSpan(
                        ordinal,
                        locator,
                        match.group(1).rstrip(".") if match else None,
                        text,
                        normalize(text),
                        "xlsx_row",
                    )
                )
        workbook.close()
    except Exception as exc:
        raise IngestionError(f"Не удалось прочитать XLSX: {exc}") from exc
    if not spans:
        warnings.append("Книга не содержит непустых ячеек")
    return spans, warnings


def parse_file(path: Path, display_name: str | None = None) -> ParsedDocument:
    validate_file(path)
    suffix = path.suffix.lower()
    parser = {".docx": parse_docx, ".pdf": parse_pdf, ".xlsx": parse_xlsx}[suffix]
    spans, warnings = parser(path)
    filename = display_name or path.name
    status = "ready" if spans and not any("заблокирован" in item for item in warnings) else "needs_attention"
    return ParsedDocument(
        filename=filename,
        file_type=suffix.lstrip("."),
        sha256=sha256_file(path),
        revision=infer_revision(filename, spans),
        status=status,
        spans=spans,
        warnings=warnings,
    )
